"""P5a documents parsing tests."""
import pytest

from app.documents import (
    ParseResult,
    UnsupportedFormatError,
    normalize_pdf_text,
    parse_document,
    strip_latex,
)


def test_strip_latex():
    tex = (
        "% a comment\n"
        "\\documentclass{article}\n"
        "\\begin{document}\n"
        "\\section{个人项目}\n"
        "\\item 负责支付平台\\textbf{架构}设计\n"
        "\\href{https://example.com}{项目链接}\n"
        "\\end{document}\n"
    )
    text = strip_latex(tex)
    assert "个人项目" in text
    assert "负责支付平台架构设计" in text
    assert "项目链接" in text
    assert "documentclass" not in text
    assert "begin" not in text
    assert "{" not in text and "}" not in text


def test_normalize_pdf_text():
    assert normalize_pdf_text("a\t  b \n\n\n c") == "a b\n\nc"


def test_parse_txt_and_md():
    assert parse_document("resume.txt", "hello\nworld".encode()).text == "hello\nworld"
    assert parse_document("resume.txt", b"x").format == "txt"
    assert parse_document("resume.md", b"# hi").format == "md"


def test_parse_tex():
    result = parse_document("resume.tex", "\\section{技能}\\item Go".encode())
    assert result.format == "tex"
    assert result.text == "技能- Go"


def test_unsupported_format():
    with pytest.raises(UnsupportedFormatError):
        parse_document("resume.xyz", b"data")


def test_parse_pdf_text_and_pages():
    import pymupdf

    doc = pymupdf.open()
    doc.new_page().insert_text((72, 72), "岗位职责：负责高并发支付平台架构设计", fontname="china-s")
    doc.new_page().insert_text((72, 72), "第二页内容", fontname="china-s")
    data = doc.tobytes()
    doc.close()

    result = parse_document("resume.pdf", data)
    assert result.format == "pdf"
    assert result.pages == 2
    assert result.page_images == []
    assert "负责高并发支付平台架构设计" in result.text
    assert "第二页内容" in result.text


def test_parse_pdf_render_pages():
    import pymupdf

    doc = pymupdf.open()
    doc.new_page().insert_text((72, 72), "简历内容", fontname="china-s")
    data = doc.tobytes()
    doc.close()

    result = parse_document("resume.pdf", data, render_pages=True)
    assert len(result.page_images) == 1
    assert result.page_images[0].startswith("data:image/jpeg;base64,")


def test_parse_docx():
    from docx import Document

    document = Document()
    document.add_paragraph("张三｜后端工程师")
    document.add_paragraph("负责支付平台架构设计")
    buffer = __import__("io").BytesIO()
    document.save(buffer)

    result = parse_document("resume.docx", buffer.getvalue())
    assert result.format == "docx"
    assert result.pages is None
    assert "张三" in result.text
    assert "负责支付平台架构设计" in result.text
