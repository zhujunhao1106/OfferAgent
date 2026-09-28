"""Document parsing (PDF/DOCX/TEX/TXT/MD) mirroring the BFF parse-pdf route.

Reference: web/src/app/api/parse-pdf/route.ts (mammoth + pdfjs-dist). The PDF
path is replaced with PyMuPDF and DOCX with python-docx per the rewrite plan.
"""
from __future__ import annotations

import base64
import io
import re
from dataclasses import dataclass, field


class UnsupportedFormatError(ValueError):
    pass


@dataclass
class ParseResult:
    text: str
    format: str
    pages: int | None = None
    page_images: list[str] = field(default_factory=list)

    def to_dict(self) -> dict:
        return {
            "text": self.text,
            "pages": self.pages,
            "pageImages": self.page_images,
            "format": self.format,
        }


def parse_document(filename: str, data: bytes, render_pages: bool = False) -> ParseResult:
    name = filename.lower()
    if name.endswith(".pdf"):
        text, pages, images = _extract_pdf(data, render_pages)
        return ParseResult(text=text, pages=pages, page_images=images, format="pdf")
    if name.endswith(".docx") or name.endswith(".doc"):
        return ParseResult(text=_extract_docx(data), format="docx")
    if name.endswith(".tex"):
        return ParseResult(text=strip_latex(data.decode("utf-8")), format="tex")
    if name.endswith(".txt") or name.endswith(".md"):
        return ParseResult(text=data.decode("utf-8"), format="md" if name.endswith(".md") else "txt")
    raise UnsupportedFormatError("Unsupported file format")


def _extract_pdf(data: bytes, render_pages: bool) -> tuple[str, int, list[str]]:
    import pymupdf

    document = pymupdf.open(stream=data, filetype="pdf")
    try:
        pages: list[str] = []
        images: list[str] = []
        rendered_limit = min(document.page_count, 3)
        for index in range(document.page_count):
            page = document.load_page(index)
            pages.append(page.get_text())
            if render_pages and index < rendered_limit:
                images.append(_render_page_as_jpeg(page))
        return normalize_pdf_text("\n\n".join(pages)), document.page_count, images
    finally:
        document.close()


def _render_page_as_jpeg(page) -> str:
    import pymupdf

    base_width = page.rect.width
    scale = min(1.5, 1100 / base_width)
    pixmap = page.get_pixmap(matrix=pymupdf.Matrix(scale, scale), alpha=False)
    jpeg = pixmap.tobytes("jpeg")
    return "data:image/jpeg;base64," + base64.b64encode(jpeg).decode("ascii")


def _extract_docx(data: bytes) -> str:
    from docx import Document

    document = Document(io.BytesIO(data))
    parts: list[str] = []
    for paragraph in document.paragraphs:
        parts.append(paragraph.text)
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            joined = " ".join(cell for cell in cells if cell)
            if joined:
                parts.append(joined)
    return "\n".join(parts).strip()


def strip_latex(tex: str) -> str:
    text = re.sub(r"%.*", "", tex)
    text = re.sub(r"\\begin\{document\}", "", text)
    text = re.sub(r"\\end\{document\}", "", text)
    text = re.sub(
        r"\\(?:documentclass|usepackage|pagestyle|geometry|setlength|renewcommand|newcommand)"
        r"\{[^}]*\}(?:\[[^\]]*\])?(?:\{[^}]*\})*",
        "",
        text,
    )
    text = re.sub(r"\\(?:section|subsection|subsubsection|textbf|textit|emph|underline|href)\{([^}]*)\}", r"\1", text)
    text = re.sub(r"\\(?:begin|end)\{[^}]*\}", "", text)
    text = re.sub(r"\\item\s*", "- ", text)
    text = re.sub(r"\\[a-zA-Z]+\*?(?:\[[^\]]*\])?(?:\{([^}]*)\})?", r"\1", text)
    text = re.sub(r"[{}]", "", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def normalize_pdf_text(text: str) -> str:
    text = re.sub(r"[\t\f\v ]+", " ", text)
    text = re.sub(r" *\n *", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()
